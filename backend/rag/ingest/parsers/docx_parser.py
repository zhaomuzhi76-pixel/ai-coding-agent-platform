"""Word .docx 解析器。职责：段落 + 表格 → ParsedPage。
不做 .doc（老格式，需 antiword）、不切块、不索引。"""
from __future__ import annotations
import os
from typing import Optional

from backend.rag.ingest.contracts import (
    ParsedPage, SourceDocument, SourceType)

MIN_PARAGRAPH_CHARS = 1


class DocxParser:
    def supports(self, path: str) -> bool:
        return path.lower().endswith(".docx")

    def parse(self, path: str, *, title: str = "",
              meta: Optional[dict] = None) -> SourceDocument:
        if not os.path.isfile(path):
            raise FileNotFoundError(path)
        try:
            import docx  # type: ignore
        except ImportError as e:
            raise RuntimeError("未安装 python-docx") from e

        d = docx.Document(path)
        pages: list = []
        # 段落
        for i, para in enumerate(d.paragraphs):
            t = (para.text or "").strip()
            if len(t) < MIN_PARAGRAPH_CHARS:
                continue
            style = getattr(para.style, "name", "") or ""
            pages.append(ParsedPage(text=t, page_no=0, section=style))
        # 表格 → 转文本
        for ti, table in enumerate(d.tables):
            rows = []
            for row in table.rows:
                cells = [str(c.text or "").strip() for c in row.cells]
                rows.append(" | ".join(cells))
            body = "\n".join(rows).strip()
            if body:
                pages.append(ParsedPage(
                    text=body, page_no=0,
                    section=f"table_{ti + 1}"))
        stat = os.stat(path)
        return SourceDocument(
            path=os.path.abspath(path), source_type=SourceType.DOCX,
            title=title or os.path.splitext(os.path.basename(path))[0],
            pages=tuple(pages), file_size=stat.st_size,
            meta={"paragraphs": len(d.paragraphs),
                  "tables": len(d.tables), **(meta or {})},
        )


def make_docx_parser() -> DocxParser:
    return DocxParser()


__all__ = ["DocxParser", "make_docx_parser"]
