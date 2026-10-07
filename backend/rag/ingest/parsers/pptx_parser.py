"""PowerPoint .pptx 解析器。职责：每张幻灯片 → 一张 ParsedPage。
包含标题、文本框、表格、备注。不做 .ppt（老格式）、不切块。"""
from __future__ import annotations
import os
from typing import Optional

from backend.rag.ingest.contracts import (
    ParsedPage, SourceDocument, SourceType)


class PptxParser:
    def supports(self, path: str) -> bool:
        return path.lower().endswith(".pptx")

    def parse(self, path: str, *, title: str = "",
              meta: Optional[dict] = None) -> SourceDocument:
        if not os.path.isfile(path):
            raise FileNotFoundError(path)
        try:
            from pptx import Presentation  # type: ignore
        except ImportError as e:
            raise RuntimeError("未安装 python-pptx") from e

        prs = Presentation(path)
        pages: list = []
        for si, slide in enumerate(prs.slides, 1):
            texts: list = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    t = (shape.text_frame.text or "").strip()
                    if t:
                        texts.append(t)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        cells = [str(c.text or "").strip()
                                 for c in row.cells]
                        texts.append(" | ".join(cells))
            # 备注
            if slide.has_notes_slide:
                notes = getattr(slide.notes_slide, "notes_text_frame", None)
                if notes is not None:
                    nt = (notes.text or "").strip()
                    if nt:
                        texts.append(f"[备注] {nt}")
            body = "\n".join(texts).strip()
            if body:
                pages.append(ParsedPage(text=body, page_no=si,
                                         section=f"slide_{si}"))
        stat = os.stat(path)
        return SourceDocument(
            path=os.path.abspath(path), source_type=SourceType.TEXT,
            title=title or os.path.splitext(os.path.basename(path))[0],
            pages=tuple(pages), file_size=stat.st_size,
            meta={"slides": len(prs.slides), **(meta or {})},
        )


def make_pptx_parser() -> PptxParser:
    return PptxParser()


__all__ = ["PptxParser", "make_pptx_parser"]
